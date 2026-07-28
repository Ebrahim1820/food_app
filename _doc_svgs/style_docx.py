"""Post-process a pandoc-generated DOCX to apply clean table styling."""
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


HEADER_BG   = RGBColor(0x15, 0x65, 0xC0)   # #1565C0 blue
HEADER_FG   = RGBColor(0xFF, 0xFF, 0xFF)   # white
ODD_BG      = RGBColor(0xFF, 0xFF, 0xFF)   # white
EVEN_BG     = RGBColor(0xE8, 0xF0, 0xFE)  # #E8F0FE light blue
BORDER_COL  = RGBColor(0xBB, 0xDE, 0xFB)  # #BBDEFB


def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    hex_val = str(rgb)  # RGBColor.__str__ returns "RRGGBB"
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_val)
    tcPr.append(shd)


def set_cell_border(cell, sides=('top', 'bottom', 'left', 'right'), color="BBDEFB", sz=4):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in sides:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'),   'single')
        el.set(qn('w:sz'),    str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def style_tables(doc: Document):
    for table in doc.tables:
        # Remove any built-in style that fights us
        table.style = doc.styles['Table']

        for row_idx, row in enumerate(table.rows):
            is_header = row_idx == 0
            bg = HEADER_BG if is_header else (ODD_BG if row_idx % 2 == 1 else EVEN_BG)

            for col_idx, cell in enumerate(row.cells):
                set_cell_bg(cell, bg)
                set_cell_border(cell)

                for para in cell.paragraphs:
                    for run in para.runs:
                        if is_header:
                            run.font.color.rgb = HEADER_FG
                            run.font.bold      = True
                            run.font.size      = Pt(10)
                        else:
                            run.font.color.rgb = RGBColor(0x21, 0x21, 0x21)
                            run.font.bold      = (col_idx == 0)
                            run.font.size      = Pt(9.5)

                    # Cell padding via tcMar
                    tc   = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    mar  = OxmlElement('w:tcMar')
                    for side in ('top', 'bottom', 'left', 'right'):
                        s = OxmlElement(f'w:{side}')
                        val = '80' if side in ('top', 'bottom') else '120'
                        s.set(qn('w:w'),    val)
                        s.set(qn('w:type'), 'dxa')
                        mar.append(s)
                    tcPr.append(mar)


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else 'DOCUMENTATION.docx'
    doc = Document(path)
    style_tables(doc)
    doc.save(path)
    print(f'✓ Styled tables in {path}')


if __name__ == '__main__':
    main()
